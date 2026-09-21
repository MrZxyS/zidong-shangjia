import json
import os
import re
import sys

import openpyxl


SOURCE = "产品信息2.xlsx"
DEVELOPMENT = "2026开发跟踪表 智利.xlsx"
start_index = int(os.environ.get("BATCH_START", "1"))
end_index = int(os.environ.get("BATCH_END", "50"))
batch_name = f"batch-{start_index}-{end_index}"
OUTPUT_DIR = f".data/{batch_name}-images"


def clean(value):
    return str(value or "").strip()


source_book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
source_sheet = source_book.active
source_rows = [list(row) for row in source_sheet.iter_rows(min_row=2, values_only=True) if clean(row[0])][start_index - 1:end_index]
target_skus = {clean(row[0]) for row in source_rows}

development_book = openpyxl.load_workbook(DEVELOPMENT, data_only=True)
development_sheet = development_book.active
rows_by_sku = {}
for row_number in range(3, development_sheet.max_row + 1):
    sku = clean(development_sheet.cell(row_number, 7).value)
    if sku in target_skus:
        rows_by_sku.setdefault(sku, []).append(row_number)

images_by_row = {}
for image in development_sheet._images:
    anchor = getattr(image, "anchor", None)
    origin = getattr(anchor, "_from", None)
    if origin is None or origin.col + 1 != 6:
        continue
    images_by_row.setdefault(origin.row + 1, []).append(image)

os.makedirs(OUTPUT_DIR, exist_ok=True)
manifest = []
for source_index, source_row in enumerate(source_rows, start=start_index):
    sku = clean(source_row[0])
    development_rows = rows_by_sku.get(sku, [])
    entry = {
        "source_index": source_index,
        "sku": sku,
        "competitor_url": clean(source_row[1]),
        "purchase_url": clean(source_row[2]),
        "source_spec": clean(source_row[3]),
        "upc": clean(source_row[4]),
        "development_rows": development_rows,
        "images": [],
    }
    if len(development_rows) == 1:
        row_number = development_rows[0]
        entry["development"] = {
            "name": clean(development_sheet.cell(row_number, 4).value),
            "color": clean(development_sheet.cell(row_number, 11).value),
            "size": clean(development_sheet.cell(row_number, 16).value),
            "category": clean(development_sheet.cell(row_number, 14).value),
            "price_clp": development_sheet.cell(row_number, 47).value,
            "weight_kg": development_sheet.cell(row_number, 23).value,
            "length_cm": development_sheet.cell(row_number, 24).value,
            "width_cm": development_sheet.cell(row_number, 25).value,
            "height_cm": development_sheet.cell(row_number, 26).value,
            "selling_points": clean(development_sheet.cell(row_number, 57).value),
        }
        for image_index, image in enumerate(images_by_row.get(development_rows[0], []), start=1):
            extension = clean(getattr(image, "format", "png")).lower()
            if extension == "jpeg":
                extension = "jpg"
            if extension not in {"jpg", "png"}:
                continue
            filename = f"{sku}-{image_index}.{extension}"
            output_path = os.path.join(OUTPUT_DIR, filename)
            try:
                image_bytes = image._data()
            except (ValueError, OSError) as error:
                entry.setdefault("image_errors", []).append(str(error))
                continue
            with open(output_path, "wb") as output:
                output.write(image_bytes)
            entry["images"].append({
                "path": output_path.replace("\\", "/"),
                "width": image.width,
                "height": image.height,
            })
    manifest.append(entry)

with open(f".data/{batch_name}-image-manifest.json", "w", encoding="utf-8") as output:
    json.dump(manifest, output, ensure_ascii=False, indent=2)

print(json.dumps({
    "target_skus": len(target_skus),
    "exact_development_matches": sum(1 for entry in manifest if len(entry["development_rows"]) == 1),
    "skus_with_images": sum(1 for entry in manifest if entry["images"]),
    "total_images": sum(len(entry["images"]) for entry in manifest),
}, ensure_ascii=False))
