import json
from collections import Counter
import openpyxl

INFO_PATH = "产品信息-店铺二.xlsx"
DEV_PATH = "2026开发跟踪表 智利-New.xlsx"
OUTPUT_PATH = ".data/store2-source.json"

def clean(value):
    return "" if value is None else str(value).strip()

info_book = openpyxl.load_workbook(INFO_PATH, read_only=True, data_only=True)
info_sheet = info_book["Sheet1"]
info_rows = []
for row_number, values in enumerate(info_sheet.iter_rows(min_row=2, values_only=True), start=2):
    sku, competitor_url, purchase_url, source_spec, upc = values[:5]
    if clean(sku):
        info_rows.append({"source_row": row_number, "sku": clean(sku),
          "competitor_url": clean(competitor_url), "purchase_url": clean(purchase_url),
          "source_spec": clean(source_spec), "upc": clean(upc)})
info_book.close()

dev_book = openpyxl.load_workbook(DEV_PATH, read_only=True, data_only=True)
dev_sheet = dev_book["开发跟踪表LHUA"]
headers = [clean(cell.value) for cell in next(dev_sheet.iter_rows(min_row=2, max_row=2))]
dev_by_sku = {}
for dev_row, values in enumerate(dev_sheet.iter_rows(min_row=3, values_only=True), start=3):
    record = dict(zip(headers, values))
    sku = clean(record.get("SKU"))
    if sku:
        dev_by_sku[sku] = {"row": dev_row, "product_name": clean(record.get("产品名称")),
          "spu": clean(record.get("SPU")), "parent_sku": clean(record.get("父SKU")),
          "color": clean(record.get("SKC（颜色）")), "category": clean(record.get("产品类目")),
          "size": clean(record.get("尺码")), "competitor_url": clean(record.get("竞品链接")),
          "purchase_url": clean(record.get("采购链接1")), "commission": clean(record.get("佣金%")),
          "weight_kg": record.get("产品重量 (kg)"), "length_cm": record.get("长(cm)"),
          "width_cm": record.get("宽(cm)"), "height_cm": record.get("高(cm)"),
          "price_clp": record.get("平台售价 CL"), "selling_points": clean(record.get("卖点"))}
dev_book.close()

counts = Counter(row["sku"] for row in info_rows)
seen, products = set(), []
for row in info_rows:
    if row["sku"] in seen: continue
    seen.add(row["sku"])
    products.append({**row, "duplicate_source_rows": counts[row["sku"]],
      "development": dev_by_sku.get(row["sku"])})

with open(OUTPUT_PATH, "w", encoding="utf-8") as handle:
    json.dump(products, handle, ensure_ascii=False, indent=2); handle.write("\n")
print(json.dumps({"rows": len(info_rows), "unique_skus": len(products),
  "matched": sum(bool(p["development"]) for p in products),
  "missing": [p["sku"] for p in products if not p["development"]]}, ensure_ascii=False))
